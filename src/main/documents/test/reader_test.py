"""Run with Dela's installed document Python; fixtures never use user documents."""

from contextlib import redirect_stdout
from concurrent.futures import ThreadPoolExecutor
import importlib.util
import io
import json
from pathlib import Path
import posixpath
import shutil
import subprocess
import sys
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile


SCRIPT = Path(__file__).resolve().parents[2] / "agent/skills/document-reading/scripts/read.py"
sys.dont_write_bytecode = True
spec = importlib.util.spec_from_file_location("document_reader", SCRIPT)
reader = importlib.util.module_from_spec(spec)
spec.loader.exec_module(reader)


def pdf_fixture():
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b""]
    children = []
    for number in range(1, 9):
        page = len(objects) + 1
        children.append(f"{page} 0 R")
        text = f"BT /F1 14 Tf 40 700 Td (Chapter {number} unique material) Tj ET" if number < 8 else ""
        if number == 2:
            for y, row in [(650, ["Item", "Value", "Unit"]), (625, ["Alpha", "42", "kg"]), (600, ["Beta", "55", "kg"])]:
                for x, word in zip([40, 220, 400], row):
                    text += f"\nBT /F1 12 Tf {x} {y} Td ({word}) Tj ET"
        content = text.encode()
        objects.extend([
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> >> >> /Contents {page + 1} 0 R >>".encode(),
            b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        ])
    objects[1] = f"<< /Type /Pages /Count 8 /Kids [{' '.join(children)}] >>".encode()
    result = b"%PDF-1.4\n"
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(result))
        result += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(result)
    result += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets[1:]:
        result += f"{offset:010d} 00000 n \n".encode()
    return result + f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()


class ReaderTests(unittest.TestCase):
    def setUp(self):
        self.temporary = TemporaryDirectory(prefix="dela-reader-test-")
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.source = self.root / "download-without-extension"
        self.source.write_bytes(pdf_fixture())
        self.output = self.root / "output"

    def run_reader(self, source=None, *options):
        stdout = io.StringIO()
        argv = [str(SCRIPT), str(source or self.source), "--output", str(self.output), *options]
        with patch("sys.argv", argv), redirect_stdout(stdout):
            reader.main()
        return json.loads(stdout.getvalue())

    def cached_job(self, name, content=None):
        job = self.root / name
        job.mkdir()
        source = job / "input"
        source.write_bytes(content if content is not None else pdf_fixture())
        (job / "document-cache.json").write_text(json.dumps({"directory": str(self.root / "cache")}))
        self.output = job / "output"
        return source

    def test_cached_outline_survives_job_release_without_extracting_full_body(self):
        first = self.cached_job("job-first")
        with patch.dict("sys.modules", {"markitdown": None}):
            original = self.run_reader(first, "--outline")
        indexed = Path(original["outline"]).read_bytes()
        self.assertFalse(original["cache"]["hit"])
        shutil.rmtree(first.parent)
        second = self.cached_job("job-second")
        with patch.object(reader, "write_outline", side_effect=AssertionError("outline parsed twice")), patch.dict("sys.modules", {"markitdown": None, "pypdfium2": None}):
            reused = self.run_reader(second, "--outline")
        self.assertTrue(reused["cache"]["hit"])
        self.assertEqual(Path(reused["outline"]).read_bytes(), indexed)
        self.assertEqual(reused["cache"]["readerVersion"], original["cache"]["readerVersion"])
        self.assertFalse((second.parent / "output/content.md").exists())

    def test_selected_text_cache_is_isolated_across_concurrent_jobs(self):
        first = self.cached_job("job-first")
        second = self.cached_job("job-second")

        def extract(source, pages):
            completed = subprocess.run([sys.executable, str(SCRIPT), str(source), "--output", str(source.parent / "output"), "--text-pages", pages], capture_output=True, text=True, check=True)
            return json.loads(completed.stdout)

        with ThreadPoolExecutor(max_workers=2) as workers:
            jobs = [workers.submit(extract, first, "1"), workers.submit(extract, second, "2")]
            results = [job.result() for job in jobs]
        self.assertIn("Chapter 1 unique material", Path(results[0]["content"]).read_text())
        self.assertNotIn("Chapter 2 unique material", Path(results[0]["content"]).read_text())
        self.assertIn("Chapter 2 unique material", Path(results[1]["content"]).read_text())
        self.assertNotIn("Chapter 1 unique material", Path(results[1]["content"]).read_text())
        shutil.rmtree(first.parent)
        third = self.cached_job("job-third")
        with patch.dict("sys.modules", {"markitdown": None, "pypdfium2": None}):
            reused = self.run_reader(third, "--text-pages", "1")
        self.assertTrue(reused["cache"]["hit"])
        self.assertEqual([entry["page"] for entry in reused["pageIndex"]], [1])
        self.assertIn("Chapter 1 unique material", Path(reused["content"]).read_text())
        self.assertTrue(Path(results[1]["content"]).exists())

    def test_cache_invalidates_on_content_and_reader_changes_and_bounds_retention(self):
        source = self.cached_job("job-first", b"# First chapter\nOriginal content")
        first = self.run_reader(source, "--outline")
        source.write_bytes(b"# Second chapter\nChanged content")
        with patch.object(reader, "CACHE_MAX_ENTRIES", 1):
            second = self.run_reader(source, "--outline")
        self.assertFalse(second["cache"]["hit"])
        self.assertNotEqual(first["cache"]["sourceHash"], second["cache"]["sourceHash"])
        self.assertEqual(len(list((self.root / "cache").iterdir())), 1)
        changed_script = self.root / "reader-revision.py"
        changed_script.write_bytes(SCRIPT.read_bytes() + b"\n")
        with patch.object(reader, "__file__", str(changed_script)):
            revised = self.run_reader(source, "--outline")
        self.assertFalse(revised["cache"]["hit"])
        self.assertNotEqual(second["cache"]["readerVersion"], revised["cache"]["readerVersion"])

    def test_page_renders_remain_local_even_with_a_persistent_cache(self):
        source = self.cached_job("job-first")
        result = self.run_reader(source, "--text-pages", "1", "--pages", "1")
        self.assertFalse(result["cache"]["enabled"])
        self.assertTrue(Path(result["images"][0]["path"]).is_relative_to(source.parent.resolve()))
        self.assertFalse((self.root / "cache").exists())

    def test_cached_office_content_materializes_its_embedded_images_in_the_new_job(self):
        content = self.pptx().read_bytes()
        source = self.cached_job("job-first", content)
        first = self.run_reader(source, "--text-pages", "2")
        original = Path(first["content"]).read_text()
        shutil.rmtree(source.parent)
        source = self.cached_job("job-second", content)
        with patch.dict("sys.modules", {"markitdown": None}):
            reused = self.run_reader(source, "--text-pages", "2")
        self.assertTrue(reused["cache"]["hit"])
        self.assertEqual(Path(reused["content"]).read_text(), original)
        self.assertTrue(reused["images"])
        for image in reused["images"]:
            self.assertTrue(Path(image["path"]).is_file())
            self.assertTrue(Path(image["path"]).is_relative_to(source.parent.resolve()))
            self.assertIn(Path(image["path"]).name, original)

    def docx(self):
        source = self.root / "lecture.docx"
        with ZipFile(source, "w") as archive:
            archive.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
            archive.writestr("_rels/.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
            archive.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>First chapter</w:t></w:r></w:p><w:p><w:r><w:t>alpha mechanism</w:t></w:r></w:p><w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Second chapter</w:t></w:r></w:p><w:tbl><w:tr><w:tc><w:p><w:r><w:t>Beta</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>42</w:t></w:r></w:p></w:tc></w:tr></w:tbl><w:sectPr/></w:body></w:document>')
        return source

    def pptx(self):
        from PIL import Image
        from pptx import Presentation
        from pptx.util import Inches

        source = self.root / "slides.pptx"
        image = self.root / "figure.png"
        Image.new("RGB", (20, 20), "navy").save(image)
        deck = Presentation()
        for number in (1, 2):
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = f"Slide {number} title"
            slide.placeholders[1].text = f"Mechanism {number} body"
            slide.notes_slide.notes_text_frame.text = f"Teacher notes {number}"
            Image.new("RGB", (20, 20), "navy" if number == 1 else "orange").save(image)
            slide.shapes.add_picture(str(image), Inches(6), Inches(2))
            if number == 2:
                table = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(4), Inches(1)).table
                table.cell(0, 0).text = "Item"
                table.cell(0, 1).text = "Value"
                table.cell(1, 0).text = "Beta"
                table.cell(1, 1).text = "42"
        deck.save(source)
        return source

    def test_pdf_full_extraction_keeps_original_pages_tables_and_empty_page_warning(self):
        result = self.run_reader()
        content = Path(result["content"]).read_text()
        self.assertEqual(result["status"], "extracted")
        self.assertEqual(result["pageCount"], 8)
        self.assertEqual([entry["page"] for entry in result["pageIndex"]], list(range(1, 9)))
        for number in range(1, 8):
            self.assertIn(f"Chapter {number} unique material", content)
        for entry in result["pageIndex"]:
            self.assertEqual(content.splitlines()[entry["startLine"] - 1], f"## PDF文件第 {entry['page']} 页")
        self.assertIn("Alpha", content)
        self.assertIn("42", content)
        self.assertIn("|", content)
        self.assertIn({"code": "pages_without_text", "pages": [8], "message": "These pages have no extracted text; inspect their page images before claiming coverage."}, result["warnings"])

    def test_pdf_page_selection_and_existing_render_option(self):
        full = self.run_reader()
        original = Path(full["content"]).read_bytes()
        result = self.run_reader(None, "--text-pages", "2-3", "--pages", "3")
        content = Path(result["content"]).read_text()
        self.assertNotIn("Chapter 1 unique material", content)
        self.assertIn("Chapter 2 unique material", content)
        self.assertIn("Chapter 3 unique material", content)
        self.assertEqual([entry["page"] for entry in result["images"]], [3])
        self.assertEqual(Path(full["content"]).read_bytes(), original)

    def test_render_only_never_imports_markitdown_or_changes_existing_text(self):
        full = self.run_reader()
        path = Path(full["content"])
        original = (path.read_bytes(), path.stat().st_mtime_ns)
        with patch.dict("sys.modules", {"markitdown": None}):
            result = self.run_reader(None, "--render-only", "--pages", "1-8")
        self.assertEqual(result["status"], "rendered")
        self.assertIsNone(result["content"])
        self.assertEqual((path.read_bytes(), path.stat().st_mtime_ns), original)
        self.assertEqual([entry["page"] for entry in result["images"]], [1, 2, 3, 4, 5, 6])
        self.assertIn({"code": "remaining_pages", "pages": [7, 8]}, result["warnings"])
        with patch("sys.stderr", io.StringIO()), self.assertRaises(SystemExit):
            self.run_reader(None, "--render-only")

    def test_pdf_outline_uses_local_index_without_body_conversion(self):
        with patch.dict("sys.modules", {"markitdown": None}):
            result = self.run_reader(None, "--outline")
        outline = json.loads(Path(result["outline"]).read_text())
        self.assertEqual(result["status"], "indexed")
        self.assertEqual(outline["unitCount"], 8)
        self.assertNotIn("text", outline["units"][0])
        self.assertEqual(outline["units"][0]["page"], 1)
        self.assertIn("Chapter 1", outline["units"][0]["titleCandidate"])
        self.assertFalse((self.output / "content.md").exists())

    def test_scanned_pdf_keeps_visual_fallback_despite_page_headings(self):
        from PIL import Image

        source = self.root / "scanned.pdf"
        Image.new("RGB", (100, 100), "navy").save(source)
        result = self.run_reader(source)
        self.assertEqual(result["status"], "needs_visual_read")
        self.assertEqual(result["pageIndex"][0]["characters"], 0)
        self.assertEqual([image["page"] for image in result["images"]], [1])
        self.assertTrue(Path(result["images"][0]["path"]).exists())

    def test_pptx_outline_and_selected_slides_preserve_notes_table_and_image(self):
        source = self.pptx()
        with patch.dict("sys.modules", {"markitdown": None}):
            indexed = self.run_reader(source, "--outline")
        outline = json.loads(Path(indexed["outline"]).read_text())
        self.assertEqual(outline["unitCount"], 2)
        self.assertEqual(outline["units"][1]["titleCandidate"], "Slide 2 title")
        result = self.run_reader(source, "--text-pages", "2")
        content = Path(result["content"]).read_text()
        self.assertNotIn("Slide 1 title", content)
        for expected in ("<!-- Slide number: 2 -->", "Slide 2 title", "Teacher notes 2", "Beta", "42"):
            self.assertIn(expected, content)
        self.assertTrue(result["images"])

    def test_docx_outline_and_section_selection_preserve_table(self):
        source = self.docx()
        with patch.dict("sys.modules", {"markitdown": None}):
            indexed = self.run_reader(source, "--outline")
        outline = json.loads(Path(indexed["outline"]).read_text())
        self.assertEqual(outline["unitCount"], 2)
        result = self.run_reader(source, "--sections", "2")
        content = Path(result["content"]).read_text()
        self.assertNotIn("First chapter", content)
        self.assertIn("Second chapter", content)
        self.assertIn("Beta", content)
        self.assertIn("42", content)
        self.assertIn("|", content)

    def test_chapter_images_never_overwrite_previous_extraction(self):
        source = self.pptx()
        first = self.run_reader(source, "--text-pages", "1")
        first_image = Path(first["images"][0]["path"])
        original = first_image.read_bytes()
        first_content = Path(first["content"]).read_text()
        second = self.run_reader(source, "--text-pages", "2")
        second_image = Path(second["images"][0]["path"])
        self.assertNotEqual(first_image, second_image)
        self.assertEqual(first_image.read_bytes(), original)
        self.assertNotEqual(second_image.read_bytes(), original)
        self.assertIn(first_image.name, first_content)
        repeat = self.run_reader(source, "--text-pages", "1")
        self.assertEqual(Path(repeat["images"][0]["path"]), first_image)

    def test_transcript_outline_keeps_timestamps_and_source_line_ranges(self):
        source = self.root / "lecture.txt"
        source.write_text("00:00 Introduction\nalpha mechanism\n00:30 Second topic\nbeta example\n", encoding="utf-8")
        with patch.dict("sys.modules", {"markitdown": None}):
            indexed = self.run_reader(source, "--outline")
        outline = json.loads(Path(indexed["outline"]).read_text())
        self.assertEqual(outline["units"][1]["timestamp"], "00:30")
        self.assertEqual(outline["units"][1]["startLine"], 3)
        result = self.run_reader(source, "--sections", "2")
        content = Path(result["content"]).read_text()
        self.assertNotIn("alpha mechanism", content)
        self.assertIn("00:30 Second topic", content)
        self.assertIn("beta example", content)

    def test_single_line_transcript_is_bounded_and_character_ranges_round_trip(self):
        source = self.root / "unbroken-lecture.txt"
        original = "00:00 " + "机制😀" * 9000
        source.write_text(original, encoding="utf-8")
        indexed = self.run_reader(source, "--outline")
        outline = json.loads(Path(indexed["outline"]).read_text())
        self.assertEqual(outline["unitCount"], 4)
        combined = []
        for entry in outline["units"]:
            self.assertEqual((entry["startLine"], entry["endLine"]), (1, 1))
            self.assertLessEqual(entry["characters"], 8000)
            expected = original[entry["charStart"]:entry["charEnd"]]
            result = self.run_reader(source, "--sections", str(entry["number"]))
            self.assertEqual(result["selectedUnits"][0]["charStart"], entry["charStart"])
            content = Path(result["content"]).read_text().split("\n\n", 1)[1]
            self.assertEqual(content, expected)
            combined.append(content)
        self.assertEqual("".join(combined), original)

    def test_docx_based_on_heading_styles_keep_chapter_boundaries(self):
        original = self.docx()
        source = self.root / "inherited-headings.docx"
        with ZipFile(original) as archive, ZipFile(source, "w") as output:
            for item in archive.infolist():
                data = archive.read(item)
                if item.filename == "word/document.xml":
                    data = data.replace(b'w:val="Heading1"', b'w:val="LectureHeading"')
                output.writestr(item, data)
            output.writestr("word/styles.xml", '<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:style w:type="paragraph" w:styleId="LectureHeading"><w:basedOn w:val="LectureBase"/></w:style><w:style w:type="paragraph" w:styleId="LectureBase"><w:basedOn w:val="Heading1"/></w:style><w:style w:type="paragraph" w:styleId="Heading1"><w:pPr><w:outlineLvl w:val="0"/></w:pPr></w:style></w:styles>')
        indexed = self.run_reader(source, "--outline")
        outline = json.loads(Path(indexed["outline"]).read_text())
        self.assertEqual(outline["unitCount"], 2)
        self.assertEqual([entry["titleCandidate"] for entry in outline["units"]], ["First chapter", "Second chapter"])
        result = self.run_reader(source, "--sections", "2")
        content = Path(result["content"]).read_text()
        self.assertNotIn("First chapter", content)
        self.assertIn("Second chapter", content)
        self.assertIn("|", content)

    def test_docx_long_paragraph_slices_preserve_run_formatting_and_exact_characters(self):
        original = self.docx()
        source = self.root / "long-paragraph.docx"
        first = "alpha " * 2000
        second = "beta " * 2000
        text = first + second
        document = f'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:rPr><w:b/></w:rPr><w:t xml:space="preserve">{first}</w:t></w:r><w:r><w:rPr><w:i/></w:rPr><w:t xml:space="preserve">{second}</w:t></w:r></w:p><w:sectPr/></w:body></w:document>'
        with ZipFile(original) as archive, ZipFile(source, "w") as output:
            for item in archive.infolist():
                output.writestr(item, document if item.filename == "word/document.xml" else archive.read(item))
        indexed = self.run_reader(source, "--outline")
        outline = json.loads(Path(indexed["outline"]).read_text())
        units = reader.structured_units(source, "docx")
        self.assertEqual(outline["unitCount"], 3)
        combined = []
        for entry in outline["units"]:
            self.assertEqual((entry["blockStart"], entry["blockEnd"]), (1, 1))
            self.assertLessEqual(entry["characters"], 8000)
            expected = text[entry["charStart"]:entry["charEnd"]]
            with reader.selected_office_stream(source, "docx", units, [entry["number"]]) as stream, ZipFile(stream) as archive:
                root = ET.fromstring(archive.read("word/document.xml"))
                actual = "".join(node.text or "" for node in root.iter(f"{reader.W}t"))
                self.assertEqual(actual, expected)
                self.assertIsNotNone(root.find(f".//{reader.W}b"))
                self.assertIsNotNone(root.find(f".//{reader.W}i"))
                combined.append(actual)
            result = self.run_reader(source, "--sections", str(entry["number"]))
            content = Path(result["content"]).read_text()
            self.assertLess(len(content), 8100)
        self.assertEqual("".join(combined), text)
        with reader.selected_office_stream(source, "docx", units, [1, 2, 3]) as stream, ZipFile(stream) as archive:
            root = ET.fromstring(archive.read("word/document.xml"))
            self.assertEqual(len(root.findall(f"{reader.W}body/{reader.W}p")), 1)
            self.assertEqual("".join(node.text or "" for node in root.iter(f"{reader.W}t")), text)

    def test_docx_character_offsets_count_run_tabs_without_changing_paragraph_tab_stops(self):
        paragraph = ET.fromstring('<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:pPr><w:tabs><w:tab w:val="left" w:pos="720"/></w:tabs></w:pPr><w:r><w:rPr><w:b/></w:rPr><w:t>甲</w:t><w:tab/><w:t>乙</w:t><w:br/><w:t>丙</w:t></w:r></w:p>')
        self.assertEqual(reader.paragraph_text(paragraph), "甲\t乙\n丙")
        selected = reader.paragraph_slice(paragraph, 2, 5)
        self.assertEqual(reader.paragraph_text(selected), "乙\n丙")
        self.assertIsNotNone(selected.find(f"{reader.W}pPr/{reader.W}tabs/{reader.W}tab"))
        self.assertIsNotNone(selected.find(f"{reader.W}r/{reader.W}rPr/{reader.W}b"))

    def test_oversized_docx_table_is_retained_intact_with_explicit_warning(self):
        original = self.docx()
        source = self.root / "large-table.docx"
        cell = "mechanism " * 1000
        with ZipFile(original) as archive, ZipFile(source, "w") as output:
            for item in archive.infolist():
                data = archive.read(item)
                if item.filename == "word/document.xml":
                    data = data.replace(b"Beta", cell.encode())
                output.writestr(item, data)
        indexed = self.run_reader(source, "--outline")
        outline = json.loads(Path(indexed["outline"]).read_text())
        large = next(entry for entry in outline["units"] if entry["characters"] > 8000)
        self.assertNotIn("charStart", large)
        self.assertTrue(any(warning["code"] == "oversized_structural_block" for warning in indexed["warnings"]))
        result = self.run_reader(source, "--sections", str(large["number"]))
        content = Path(result["content"]).read_text()
        self.assertIn(cell.strip(), content)
        self.assertIn("42", content)
        self.assertIn("|", content)
        self.assertTrue(any(warning["code"] == "oversized_structural_block" for warning in result["warnings"]))

    def test_pptx_absolute_package_notes_relationship_is_readable(self):
        original = self.pptx()
        source = self.root / "absolute-notes.pptx"
        with ZipFile(original) as archive, ZipFile(source, "w") as output:
            for item in archive.infolist():
                data = archive.read(item)
                if item.filename.startswith("ppt/slides/_rels/"):
                    relationships = ET.fromstring(data)
                    for relationship in relationships:
                        if (relationship.get("Type") or "").endswith("/notesSlide"):
                            relationship.set("Target", "/" + posixpath.normpath(posixpath.join("ppt/slides", relationship.get("Target"))))
                    data = ET.tostring(relationships, encoding="utf-8", xml_declaration=True)
                output.writestr(item, data)
        indexed = self.run_reader(source, "--outline")
        self.assertEqual(indexed["unitCount"], 2)
        result = self.run_reader(source, "--text-pages", "2")
        content = Path(result["content"]).read_text()
        self.assertIn("Teacher notes 2", content)

    def test_non_pdf_default_conversion_still_extracts_full_source(self):
        for source in (self.docx(), self.pptx()):
            download = self.root / "office-download-without-extension"
            download.write_bytes(source.read_bytes())
            result = self.run_reader(download)
            content = Path(result["content"]).read_text()
            self.assertEqual(result["status"], "extracted")
            self.assertIn("First chapter" if source.suffix == ".docx" else "Slide 1 title", content)
            self.assertIn("Second chapter" if source.suffix == ".docx" else "Slide 2 title", content)


if __name__ == "__main__":
    unittest.main()
